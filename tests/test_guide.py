# -*- coding: utf-8 -*-
"""参考答案指引的纯逻辑测试：分档、指引串、表格骨架。

不起服务、不查库——这几样一旦错了，表和网页会同时错，而且错得很安静
（比如分档边界挪一格，表里所有"中"都变成"低"，没人会注意到）。
"""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from src.guide import (  # noqa: E402
    DETAIL_HEAD, band, build_workbook, clean_title, guess_company, guide_of,
    make_meta, round_cn,
)


def test_band_边界():
    assert band(0.90) == "高"
    assert band(0.60) == "高"
    assert band(0.599) == "中"
    assert band(0.42) == "中"
    assert band(0.419) == "低"
    assert band(0.30) == "低"
    assert band(0.299) == "极低"
    assert band(0.0) == "极低"


def test_轮次与标题():
    assert round_cn(1) == "一轮回复"
    assert round_cn("2") == "二轮回复"
    assert round_cn(None) == "第None轮回复"
    # 标题里自带的编号要去掉，表格另有「题号」列
    assert clean_title("问题3.关于收入确认") == "关于收入确认"
    assert clean_title("12、关于毛利率") == "关于毛利率"


def test_公司名从文件名抠():
    name = "关于深圳市裕富照明股份有限公司公开发行股票并在北交所上市申请文件的审核问询函.docx"
    assert guess_company(name) == "深圳市裕富照明股份有限公司"
    assert guess_company("通知.docx") == ""      # 抠不出就空着，不猜


def _hit(company, score, question="关于收入确认", board="北交所", rnd=1, no=3):
    return {"company": company, "score": score, "board": board, "round": rnd,
            "question": question, "question_no": no, "raw_text": "说明" + company + "的收入确认政策",
            "question_id": company + "-q3"}


def _rows():
    return [
        {"qno": 1, "qtitle": "股权结构", "sub_idx": 1, "ask": "说明股权代持情况",
         "hits": [_hit("甲公司", 0.71), _hit("乙公司", 0.44), _hit("丙公司", 0.31)],
         "best_any": 0.71},
        # 这个小问一条都没命中：表里要占一行写「未找到」，不能让行号跳过去
        {"qno": 2, "qtitle": "收入确认", "sub_idx": 1, "ask": "说明收入确认政策",
         "hits": [], "best_any": 0.12},
        {"qno": 2, "qtitle": "收入确认", "sub_idx": 2, "ask": "说明境外收入的真实性",
         "hits": [_hit("甲公司", 0.55)], "best_any": 0.55},
    ]


def test_指引串():
    assert guide_of(_hit("甲公司", 0.7)) == "北交所/甲公司/一轮回复/关于收入确认"


def test_表格骨架():
    rows = _rows()
    meta = {"ts": "2026-10-29 16:00", "source": "某函.docx", "n_q": 2, "n_sub": 2,
            "corpus": "1 份函 / 1 条提问 / 6 个板块", "top": 5, "min_score": 0.0}
    wb = build_workbook(rows, meta, 5)
    assert wb.sheetnames == ["说明", "命中排行", "明细TOP2", "明细TOP5"]

    top2 = wb["明细TOP2"]
    assert [c.value for c in top2[1]] == DETAIL_HEAD          # 8 列，一列不多
    # 题1 留 2 条 + 题2小问1「未找到」1 行 + 题2小问2 留 1 条 + 表头
    assert top2.max_row == 5
    assert top2.cell(row=2, column=5).value == "甲公司"
    assert top2.cell(row=3, column=5).value == "乙公司"       # TOP2 里第三名不该出现
    assert [c.value for c in top2[4]][:5] == [2, "收入确认", 1, "说明收入确认政策", "未找到"]

    top5 = wb["明细TOP5"]
    assert top5.max_row == 6                                  # 3 条 + 1 行未找到 + 1 条 + 表头
    # 相似度只写档位，不写数值
    assert [top5.cell(row=r, column=7).value for r in (2, 3, 4)] == ["高", "中", "低"]

    rank = wb["命中排行"]
    assert rank.cell(row=2, column=1).value == "北交所/甲公司/一轮回复"
    assert rank.cell(row=2, column=2).value == 2              # 两个小问都撞上，算 2


def test_表能落盘(tmp_path=None):
    import tempfile

    rows = _rows()
    meta = {"ts": "2026-10-29 16:00", "source": "某函.docx", "n_q": 2, "n_sub": 2,
            "corpus": "1 份函", "top": 2, "min_score": 0.0}
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "t.xlsx"
        build_workbook(rows, meta, 2).save(str(p))
        assert p.exists() and p.stat().st_size > 3000
        assert load_workbook(str(p)).sheetnames[0] == "说明"


if __name__ == "__main__":
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print("ok  " + name)
    print("全部通过")
