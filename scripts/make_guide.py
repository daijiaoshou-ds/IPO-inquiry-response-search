# -*- coding: utf-8 -*-
"""把一份问询函做成「参考答案指引」：拆题 → 逐小问检索 → 汇总成表。

给人看的，不是给 AI 看的。表里只有"去哪看"：哪个板块、哪家公司、第几轮、
第几题、第几页，以及撞上的是哪一句问法。不做任何答案总结——答案要人自己翻。

用法：
    .venv\\Scripts\\python.exe -m scripts.make_guide --pdf 新收到的函.pdf
    .venv\\Scripts\\python.exe -m scripts.make_guide --pdf xxx.pdf --out 指引.xlsx --top 5
    .venv\\Scripts\\python.exe -m scripts.make_guide --text 粘贴的题.txt --md

输出（默认到 out/ 目录）：
    <函名>_参考答案指引_<日期>.xlsx    三张表：指引 / 明细 / 说明
    <函名>_参考答案指引_<日期>.md      同样的内容，纯文本版（--md 时生成）
"""
from __future__ import annotations

import argparse
import re
import sqlite3
import sys
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# 分档：实测"库里没有的主题" top1 落在 0.22~0.37，真相近的在 0.42~0.72，
# 语料里有逐字相同提问的会顶到 0.9 以上（分句会被 0.9 折扣压一下）。
BANDS = ((0.60, "高"), (0.42, "中"), (0.30, "弱"))
DEFAULT_MIN = 0.30
# 同一家公司只留一条：r1/r2/r3 常常逐字重复同一道题，占满一整列没有信息量
MAX_PER_COMPANY = 1


def clean_title(title: str) -> str:
    """标题里已经带「问题3.」这类编号，表格另有题号列，去掉避免重复。"""
    return re.sub(r"^\s*(?:问题\s*\d{1,3}|\d{1,3})\s*[\.、]\s*", "", title or "").strip()


def band(score: float, exact: bool) -> str:
    if exact:
        return "逐字"
    for cut, name in BANDS:
        if score >= cut:
            return name
    return "弱"


# ---------------------------------------------------------------- 读入

def load_letter(path: Path, board: str):
    """PDF / docx / txt 三路统一入口，拆分规则都在 src/ingest.py。"""
    from src.ingest import load_file
    return load_file(path, board=board)


# ---------------------------------------------------------------- 检索

def aggregate(conn, ask: str, top: int, min_score: float,
              exclude: set[str]) -> tuple[list[dict], float]:
    """一个小问 -> (至多 top 条历史相似题, 全部候选里的最高分)。

    最高分单独返回：门槛以下一条不给，但要能告诉人"最像的只有 0.28"，
    这跟"工具没查到"是两件事。
    """
    from src.search import search_grouped

    res = search_grouped(conn, ask, top=max(top * 8, 40))
    best_any = 0.0
    for grp in res["groups"]:
        for h in grp["hits"]:
            if h["company"] not in exclude and h["score"] > best_any:
                best_any = h["score"]
    best: "OrderedDict[str, dict]" = OrderedDict()
    for part, grp in zip(res["parts"], res["groups"]):
        for h in grp["hits"]:
            if h["score"] < min_score:
                continue
            if h["company"] in exclude:
                continue          # 别自己匹配自己
            key = h["question_id"]
            cur = best.get(key)
            if cur is None or h["score"] > cur["score"]:
                best[key] = {
                    "question_id": key,
                    "score": h["score"], "exact": h["exact"],
                    "board": h["board"], "company": h["company"],
                    "round": h["round"], "qno": h["question_no"],
                    "title": (h["question"] or "").strip(),
                    "pages": h["pages"], "via": h["raw_text"],
                    "q_id": h["q_id"],
                }

    ranked = sorted(best.values(), key=lambda x: (-x["score"], x["company"]))
    seen_text: set[str] = set()
    per_company: dict[tuple, int] = {}
    keep = []
    for x in ranked:
        ck = (x["company"], x["board"])
        if per_company.get(ck, 0) >= MAX_PER_COMPANY:
            continue
        per_company[ck] = per_company.get(ck, 0) + 1
        keep.append(x)
        if len(keep) >= top:
            break
    return keep, best_any


def file_of(conn, q_ids: list[str]) -> dict[str, str]:
    """命中的那条 query -> 原文 PDF 相对路径（方便人直接翻）。"""
    out = {}
    for i in range(0, len(q_ids), 400):
        chunk = q_ids[i:i + 400]
        ph = ",".join("?" * len(chunk))
        for qid, fp in conn.execute(
                "SELECT q.q_id, l.file_path FROM query q "
                "JOIN letter l ON l.letter_id = q.letter_id "
                "WHERE q.q_id IN (" + ph + ")", chunk):
            out[qid] = fp
    return out


# ---------------------------------------------------------------- 输出

def write_xlsx(path: Path, letter, rows: list[dict], meta: dict, top: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="EEF3FA")
    wrap = Alignment(wrap_text=True, vertical="top")

    # ---- 表1 指引：一题一块，一行一个小问
    ws = wb.active
    ws.title = "指引"
    headers = ["题号", "问题", "小问", "命中"]
    for i in range(1, top + 1):
        headers += ["参考%d 出处" % i, "参考%d 相似度" % i, "参考%d 撞上的问法" % i]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border

    r = 2
    for row in rows:
        line = [row["qno"], row["qtitle"], row["ask"], len(row["hits"])]
        for h in row["hits"]:
            line.append("%s·%s·r%s·题%s·p%s" % (
                h["board"], h["company"], h["round"], h["qno"], h["pages"]))
            line.append("%s %.2f" % (band(h["score"], h["exact"]), h["score"]))
            line.append(h["via"][:90])
        if not row["hits"]:
            line.append("—（最高 %.2f，未达门槛）" % row.get("best_any", 0.0))
        while len(line) < len(headers):
            line.append("")
        ws.append(line)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = wrap
            cell.border = border
            cell.font = Font(size=10)
        r += 1

    widths = [6, 26, 42, 6]
    for _ in range(top):
        widths += [30, 10, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"

    # ---- 表2 明细：一行一条参考，方便筛选排序
    ws2 = wb.create_sheet("明细")
    h2 = ["题号", "问题", "小问", "排名", "相似度档", "分数", "板块", "公司",
          "轮次", "历史题号", "该题标题", "页码", "撞上的问法", "原文文件"]
    ws2.append(h2)
    for c in range(1, len(h2) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.font = Font(bold=True, size=10)
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    r = 2
    for row in rows:
        for i, h in enumerate(row["hits"], 1):
            ws2.append([row["qno"], row["qtitle"], row["ask"], i,
                        band(h["score"], h["exact"]), round(h["score"], 3),
                        h["board"], h["company"], "r%s" % h["round"], h["qno"],
                        clean_title(h["title"]), h["pages"], h["via"],
                        h.get("file", "")])
            for c in range(1, len(h2) + 1):
                cell = ws2.cell(row=r, column=c)
                cell.alignment = wrap
                cell.border = border
                cell.font = Font(size=10)
            r += 1
    for i, w in enumerate([6, 26, 42, 6, 10, 8, 8, 20, 8, 8, 30, 10, 46, 44], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "D2"
    if r > 2:
        ws2.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(h2)), r - 1)

    # ---- 表3 说明
    ws3 = wb.create_sheet("说明")
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 96
    lines = [
        ("生成时间", meta["ts"]),
        ("待答函", meta["source"]),
        ("拆出", "%d 道题 / %d 个小问" % (meta["n_q"], meta["n_sub"])),
        ("语料覆盖", meta["corpus"]),
        ("检索参数", "每个小问取 %d 条；相似度低于 %.2f 丢弃；同一家公司只留 1 条" % (top, meta["min_score"])),
        ("", ""),
        ("相似度档", "高 ≥0.60 ／ 中 0.42–0.60 ／ 弱 0.30–0.42 ／ 逐字 = 语料里有逐字相同的提问"),
        ("怎么用", "按档从高往低看。档位只说明\"问法像不像\"，不说明答案对不对，"
                 "答案一律翻原文——原文路径在「明细」最后一列，页码在第 12 列。"),
        ("为什么没答案", "本表只做定位，不做答案摘要。问询回复要严谨，答案必须人自己看。"),
        ("", ""),
        ("注意", "相似度是 4-gram 逐字比对，不是语义。"
                 "\"存货跌价准备计提是否充分\"和\"转回是否合理\"在向量空间很近，"
                 "在监管逻辑里是两回事，所以宁可给弱档也不替你判断。"),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
        c = ws3.cell(row=i, column=2, value=v)
        c.font = Font(size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(str(path))


def write_md(path: Path, rows: list[dict], meta: dict) -> None:
    L = ["# 参考答案指引", "",
         "- 待答函：%s" % meta["source"],
         "- 拆出：%d 道题 / %d 个小问" % (meta["n_q"], meta["n_sub"]),
         "- 语料：%s" % meta["corpus"],
         "- 生成：%s" % meta["ts"], "",
         "> 只给定位，不给答案摘要。相似度分档：高 ≥0.60 / 中 0.42–0.60 / "
         "弱 0.30–0.42 / 逐字 = 语料里有逐字相同的提问。", ""]

    by_q: "OrderedDict[tuple, list[dict]]" = OrderedDict()
    for row in rows:
        by_q.setdefault((row["qno"], row["qtitle"]), []).append(row)

    for (qno, qtitle), subs in by_q.items():
        L.append("## 题%s　%s" % (qno, qtitle or "—"))
        L.append("")
        for si, row in enumerate(subs, 1):
            L.append("**小问 %d**：%s" % (si, row["ask"]))
            L.append("")
            if not row["hits"]:
                L.append("_语料里没有相近的提问（最像的一条只有 %.2f，未达门槛）。_"
                         % row.get("best_any", 0.0))
                L.append("")
                continue
            L.append("| # | 相似度 | 出处 | 库里那家的问法 |")
            L.append("| --- | --- | --- | --- |")
            for i, h in enumerate(row["hits"], 1):
                L.append("| %d | %s %.2f | %s·%s·r%s·题%s·p%s | %s |" % (
                    i, band(h["score"], h["exact"]), h["score"], h["board"],
                    h["company"], h["round"], h["qno"], h["pages"],
                    h["via"][:70].replace("|", "／")))
            L.append("")
    path.write_text("\n".join(L), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--pdf", help="待答的问询函 PDF")
    src.add_argument("--text", help="待答的题目文本文件（.txt/.md）")
    src.add_argument("--docx", help="待答的问询函 Word（.docx）")
    ap.add_argument("--out", help="输出 xlsx 路径；不给就放 out/")
    ap.add_argument("--board", default="问询函",
                    help="板块标签，只影响抽取器走哪套规则。新收到的问询函用默认值就行")
    ap.add_argument("--top", type=int, default=5, help="每个小问给几条参考（默认 5）")
    ap.add_argument("--min-score", type=float, default=DEFAULT_MIN)
    ap.add_argument("--exclude", default=None,
                    help="排除这家公司（默认按文件名自动排除，避免自己匹配自己）")
    ap.add_argument("--md", action="store_true", help="同时输出 Markdown 版")
    args = ap.parse_args()

    t0 = time.perf_counter()
    if args.pdf:
        path = Path(args.pdf)
    elif args.docx:
        path = Path(args.docx)
    else:
        path = Path(args.text)
    if not path.exists():
        raise SystemExit("找不到文件：" + str(path))
    try:
        letter = load_letter(path, args.board)
    except (ValueError, FileNotFoundError) as e:
        raise SystemExit(str(e))
    source = path.name
    t_split = time.perf_counter() - t0

    # 自动排除自己：文件名里的公司名，若库里存在就排掉
    exclude = set()
    if args.exclude:
        exclude.add(args.exclude)
    else:
        stem = re.sub(r"_r\d+$", "", Path(source).stem)
        exclude.add(stem)
    conn = sqlite3.connect(str(ROOT / "data" / "queries.db"))
    conn.row_factory = sqlite3.Row

    qmap = {q["no"]: q for q in letter.questions}
    rows = []
    t1 = time.perf_counter()
    for q in letter.queries:
        hits, best_any = aggregate(conn, q.raw_text, args.top, args.min_score, exclude)
        rows.append({
            "qno": q.question_no,
            "qtitle": clean_title(qmap.get(q.question_no, {}).get("title", "")),
            "ask": q.raw_text,
            "hits": hits,
            "best_any": best_any,
        })
    t_search = time.perf_counter() - t1

    # 补原文路径
    all_qids = list({h["q_id"] for r in rows for h in r["hits"]})
    fmap = file_of(conn, all_qids)
    for r in rows:
        for h in r["hits"]:
            h["file"] = fmap.get(h["q_id"], "")

    st = conn.execute("SELECT COUNT(*), (SELECT COUNT(*) FROM query) FROM letter").fetchone()
    meta = {
        "ts": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "source": source,
        "n_q": len(letter.questions),
        "n_sub": len(letter.queries),
        "corpus": "%d 份函 / %d 条提问 / 6 个板块（截至生成时）" % (st[0], st[1]),
        "min_score": args.min_score,
    }

    out_dir = ROOT / "out"
    out_dir.mkdir(exist_ok=True)
    stem = re.sub(r"\.(pdf|txt|md|docx)$", "", source, flags=re.I)
    date = datetime.now().strftime("%Y%m%d")
    xlsx = Path(args.out) if args.out else out_dir / ("%s_参考答案指引_%s.xlsx" % (stem, date))
    write_xlsx(xlsx, letter, rows, meta, args.top)
    made = [xlsx]
    if args.md:
        md = xlsx.with_suffix(".md")
        write_md(md, rows, meta)
        made.append(md)

    total_hits = sum(len(r["hits"]) for r in rows)
    empty = sum(1 for r in rows if not r["hits"])
    print("拆题 %.1f s ／ 检索 %.1f s ／ 合计 %.1f s" % (
        t_split, t_search, time.perf_counter() - t0))
    print("拆出 %d 道题 / %d 个小问；取回 %d 条参考（%d 个小问没到门槛）" % (
        len(letter.questions), len(letter.queries), total_hits, empty))
    for m in made:
        print("已生成：" + str(m))


if __name__ == "__main__":
    main()
